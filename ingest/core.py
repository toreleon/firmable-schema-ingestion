from __future__ import annotations

import csv
import hashlib
import io
from itertools import islice
import json
import re
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
import xml.etree.ElementTree as ET

import yaml

VERSION = '0.1.0'
TRANSFORMS = {'strip', 'upper', 'lower', 'digits', 'abn', 'acn', 'postcode', 'date', 'enum', 'split', 'regex'}


def now():
    return datetime.now(timezone.utc).isoformat()


def digest(value):
    raw = value if isinstance(value, bytes) else json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()
    return hashlib.sha256(raw).hexdigest()


def config_digest(config):
    return digest({k:v for k,v in config.items() if k != 'review'})


def record_id(source,index,row):
    return f"{source['resource_sha256'][:16]}:{index}:{digest(row)[:16]}"


def inventory(rows,source,config,observations):
    counts = {}
    for obs in observations:
        parent = obs.get('provenance',{}).get('parent_source_record_id',obs['source_record_id'])
        counts[parent] = counts.get(parent,0)+1
    return [{'source_id':source['source_id'],'source_record_id':record_id(source,i,row),
             'entity_role':config.get('entity_role','unknown'),'observation_count':counts.get(record_id(source,i,row),0),
             'row_number':i,'resource_sha256':source['resource_sha256'],
             'business_candidate_basis':'Agent classified source role; individual/business identity may remain ambiguous'}
            for i,row in enumerate(rows,1)]


def save(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + '\n')
    tmp.replace(path)


def read(path):
    return json.loads(Path(path).read_text())


def jsonl(path, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + '.tmp')
    with tmp.open('w') as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False, default=str) + '\n')
    tmp.replace(path)


def load_jsonl(path):
    with Path(path).open() as f:
        return [json.loads(line) for line in f if line.strip()]


def fetch(url, max_bytes=40_000_000, timeout=45):
    if urlparse(url).scheme not in {'http', 'https'}:
        raise ValueError('Only public HTTP(S) resources supported')
    req = urllib.request.Request(url, headers={'User-Agent': 'ReviewableIngestion/0.1 (public-data technical assignment)'})
    with urllib.request.urlopen(req, timeout=timeout) as response:
        if response.status != 200 or response.headers.get('x-amzn-waf-action'):
            raise ValueError(f'HTTP {response.status}: resource unavailable or WAF challenge')
        raw = response.read(max_bytes + 1)
        if not raw:
            raise ValueError('Empty resource response')
        if len(raw) > max_bytes:
            raise ValueError(f'Resource exceeds {max_bytes} byte cap; choose a bounded resource')
        return raw, dict(response.headers)


def getpath(row, path):
    # Exact key first: CSV headings often contain periods.
    if isinstance(row, dict) and path in row:
        return row[path]
    node = row
    for part in path.split('.'):
        if isinstance(node, dict):
            node = node.get(part)
        elif isinstance(node, list) and part.isdigit() and int(part) < len(node):
            node = node[int(part)]
        else:
            return None
    return node


def fields_of(row, prefix=''):
    out = []
    for key, val in row.items():
        path = f'{prefix}.{key}' if prefix else str(key)
        if isinstance(val, dict):
            out += fields_of(val, path)
        else:
            out.append(path)
    return out


def decode(raw):
    import codecs
    encodings = ('utf-16',) if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','cp1252')
    for encoding in encodings:
        try:
            # Prefix inspection may end within a multibyte codepoint; do not misclassify UTF-8.
            return codecs.getincrementaldecoder(encoding)().decode(raw,final=False), encoding
        except (UnicodeError, LookupError):
            pass
    raise ValueError('Unsupported encoding')


def sniff(raw):
    if raw.startswith(b'PK'):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            names = z.namelist()
            if '[Content_Types].xml' in names:
                return {'format': 'xlsx', 'sheet': 0, 'header_row': 1}
            return {'format': 'zip', 'members': names[:50]}
    text, encoding = decode(raw[:100_000])
    head = text.lstrip()
    if head.startswith(('{', '[')):
        # Two complete object lines distinguish JSONL from a pretty-printed JSON
        # document. Inspect only the prefix; a later row can be outside the sample.
        lines=list(islice((line for line in text.splitlines() if line.strip()),2))
        if len(lines)==2:
            try:
                if all(isinstance(json.loads(line),dict) for line in lines):
                    return {'format':'jsonl','encoding':encoding}
            except json.JSONDecodeError:
                pass
        return {'format': 'json', 'record_path': '', 'encoding': encoding}
    if head.startswith('<'):
        return {'format': 'xml', 'record_tag': '', 'encoding': encoding}
    try:
        delim = csv.Sniffer().sniff(text[:20000], delimiters=',;\t|').delimiter
    except csv.Error:
        delim = ','
    return {'format': 'csv', 'delimiter': delim, 'header_row': 1, 'encoding': encoding}


def parse_records(raw, parser, limit=1000):
    fmt = parser['format'].lower()
    if fmt == 'zip':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            member = parser.get('member')
            if not member or member not in z.namelist():
                raise ValueError('ZIP requires an explicit member')
            info = z.getinfo(member)
            if info.file_size > 40_000_000 or info.file_size / max(info.compress_size, 1) > 200:
                raise ValueError('Archive expansion limit exceeded')
            return parse_records(z.read(member), parser['inner'], limit)
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = parser.get('sheet', 0)
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        values = ws.iter_rows(values_only=True)
        for _ in range(int(parser.get('header_row', 1)) - 1):
            next(values)
        headings = [str(x).strip() if x is not None else f'unnamed_{i}' for i, x in enumerate(next(values))]
        if len(headings) != len(set(headings)):
            raise ValueError('Duplicate headings: agent must select the correct header row')
        out = []
        for row in values:
            if any(x is not None for x in row):
                out.append(dict(zip(headings, [x.isoformat() if hasattr(x, 'isoformat') else x for x in row])))
            if len(out) >= limit:
                break
        wb.close()
        return out
    text = raw.decode(parser.get('encoding', 'utf-8-sig'))
    if fmt == 'csv':
        lines = io.StringIO(text)
        for _ in range(int(parser.get('header_row', 1)) - 1):
            next(lines)
        reader = csv.DictReader(lines, delimiter=parser.get('delimiter', ','))
        if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError('Missing or duplicate headings')
        out = []
        for row in reader:
            if None in row:
                raise ValueError('Ragged CSV row: revise delimiter or header row')
            if any(x not in (None, '') for x in row.values()):
                out.append(row)
            if len(out) >= limit:
                break
        return out
    if fmt in {'json', 'geojson', 'jsonl'}:
        if fmt == 'jsonl':
            rows = [json.loads(line) for line in islice((line for line in io.StringIO(text) if line.strip()),limit)]
        else:
            obj = json.loads(text)
            path = parser.get('record_path', '')
            rows = getpath(obj, path) if path else obj
            if isinstance(rows, dict):
                raise ValueError('JSON object requires a record_path to an array')
        if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows[:limit]):
            raise ValueError('Expected an array of records')
        return rows[:limit]
    if fmt == 'xml':
        if '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
            raise ValueError('XML entities/DTDs are unsupported')
        tag = parser.get('record_tag')
        if not tag:
            raise ValueError('XML needs record_tag')
        root = ET.fromstring(text)
        return [{child.tag: child.text for child in elem} for elem in root.iter(tag)][:limit]
    raise ValueError(f'Unsupported format: {fmt}; route to manual adapter review')


def inspect_resource(raw):
    guess = sniff(raw)
    report = {'sha256': digest(raw), 'bytes': len(raw), 'parser_guess': guess}
    if guess['format'] == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        report['sheets'] = {ws.title: [[str(v)[:180] if v is not None else None for v in row] for row in list(__import__('itertools').islice(ws.iter_rows(values_only=True), 12))] for ws in wb.worksheets[:8]}
        wb.close()
    elif guess['format'] == 'zip':
        report['members'] = guess['members']
    else:
        text, _ = decode(raw[:30000])
        report['head'] = text[:16000]
    return report


def valid_abn(value):
    if not re.fullmatch(r'\d{11}', str(value), flags=re.ASCII):
        return False
    digits = [int(c) for c in value]
    digits[0] -= 1
    return sum(a*b for a,b in zip(digits, [10,1,3,5,7,9,11,13,15,17,19])) % 89 == 0


def valid_acn(value):
    if not re.fullmatch(r'\d{9}', str(value), flags=re.ASCII) or set(value) == {'0'}:
        return False
    return (10 - sum(int(c)*w for c,w in zip(value[:8], range(8, 0, -1))) % 10) % 10 == int(value[8])


def transform(value, steps):
    if value is None or str(value).strip().lower() in {'', 'n/a', 'null', 'none', 'not applicable', '-'}:
        return None
    if isinstance(value, (dict, list)) and not any((s if isinstance(s,str) else s.get('op')) == 'split' for s in steps):
        raise ValueError('Structured values need an explicit split mapping')
    for step in steps:
        step = {'op': step} if isinstance(step, str) else step
        op = step['op']
        if isinstance(value,list) and op not in {'split','regex'}:
            value = [transform(item,[step]) for item in value]
            continue
        if op not in TRANSFORMS:
            raise ValueError(f'Unknown transform {op}')
        if op in {'strip', 'upper', 'lower'}:
            value = getattr(str(value), op)()
        elif op in {'digits', 'abn', 'acn', 'postcode'}:
            # Spreadsheet numeric values may be integral; do not append a decimal zero.
            value = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
            if re.search(r'[A-Za-z]', value) or not re.fullmatch(r'[0-9\s.\-()]+', value):
                raise ValueError('Identifier contains non-numeric text')
            value = re.sub(r'[^0-9]', '', value)
            if op == 'abn' and not valid_abn(value):
                raise ValueError('Invalid ABN checksum')
            if op == 'acn' and not valid_acn(value):
                raise ValueError('Invalid ACN checksum')
            if op == 'postcode':
                value = value.zfill(4)
                if len(value) != 4:
                    raise ValueError('Invalid postcode length')
        elif op == 'enum':
            lookup = {str(k).casefold():v for k,v in step.get('map', {}).items()}
            value = lookup.get(str(value).casefold(), step.get('default'))
        elif op == 'date':
            if step.get('format'):
                value = datetime.strptime(str(value), step['format']).date().isoformat()
            else:
                value = datetime.fromisoformat(str(value).replace('Z', '+00:00')).date().isoformat()
        elif op == 'split':
            value = value if isinstance(value, list) else str(value).split(step.get('delimiter', ';'))
        elif op == 'regex':
            import regex
            pattern = step['pattern']
            if len(pattern) > 1000 or len(str(value)) > 20000:
                raise ValueError('Regex input/pattern limit exceeded')
            try:
                matches = list(regex.finditer(pattern, str(value), timeout=.05, flags=regex.MULTILINE))
            except TimeoutError as exc:
                raise ValueError('Regex timeout') from exc
            if step.get('all'):
                if not matches:
                    return None
                value = [match.group(step.get('group',1)) for match in matches]
                continue
            if len(matches) != 1:
                raise ValueError('Regex needs exactly one match; ambiguous compound record refused')
            value = matches[0].group(step.get('group',1))
    return value


def ontology_fields(ontology):
    return {f'{group}.{name}': spec for group in ('entity', 'address') for name, spec in ontology[group].items()}


def validate_value(value, target, ontology):
    spec = ontology_fields(ontology)[target]
    value = str(value).strip()
    if not value:
        raise ValueError('Empty canonical value')
    if spec.get('pattern') and not re.fullmatch(spec['pattern'], value, flags=re.ASCII):
        raise ValueError('Ontology pattern mismatch')
    if 'values' in spec and value not in spec['values']:
        raise ValueError('Value outside ontology enum')
    if target == 'entity.abn' and not valid_abn(value):
        raise ValueError('Invalid ABN checksum')
    if target == 'entity.acn' and not valid_acn(value):
        raise ValueError('Invalid ACN checksum')
    if spec['type'] == 'date':
        datetime.strptime(value, '%Y-%m-%d')
    if target == 'entity.website' and urlparse(value).scheme not in {'https','http'}:
        raise ValueError('Website requires an HTTP(S) URL')
    if target == 'address.postcode' and not re.fullmatch(r'\d{4}', value, flags=re.ASCII):
        raise ValueError('Australian postcode requires four digits')
    return value


def validate_config(config, rows, ontology, diagnostic=False):
    known = set().union(*(set(fields_of(row)) for row in rows)) if rows else set()
    errors = []
    if not diagnostic and config.get('minimum_field_confidence',.8) < .8:
        errors.append('minimum_field_confidence cannot be below the engine policy floor 0.8')
    targets = ontology_fields(ontology)
    for m in config.get('mappings', []):
        if m.get('source_field') not in known:
            errors.append(f"Unknown source field: {m.get('source_field')}")
        if m.get('canonical_field') not in targets:
            errors.append(f"Unknown canonical field: {m.get('canonical_field')}")
        if m.get('separate_entities') and m.get('canonical_field') not in {'entity.abn','entity.acn','entity.nzbn'}:
            errors.append('Separate entity claims are permitted only for explicit business identifiers')
        if not isinstance(m.get('confidence'), (float, int)) or not 0 <= m['confidence'] <= 1:
            errors.append('Mapping confidence outside [0,1]')
        for step in m.get('transforms', []):
            if (step if isinstance(step,str) else step.get('op')) not in TRANSFORMS:
                errors.append('Unsupported transform')
    if not config.get('mappings') and not config.get('unmapped_fields'):
        errors.append('Empty config must explain every unmapped field')
    if errors:
        raise ValueError('; '.join(errors))


def extract(rows, config, source, ontology, ingested_at=None, row_numbers=None, diagnostic=False):
    validate_config(config, rows, ontology,diagnostic=diagnostic)
    ingested_at = ingested_at or source.get('downloaded_at') or now()
    config_hash = config_digest(config)
    if config.get('ontology_sha256') and config['ontology_sha256'] != digest(ontology):
        raise ValueError('Ontology fingerprint changed; regenerate and review the mapping')
    observations, errors = [], []
    for index, row in zip(row_numbers or range(1, len(rows)+1), rows):
        raw_hash = digest(row)
        # Snapshot + parsed row position guarantees exact traceability, including duplicate rows.
        current_record_id = record_id(source,index,row)
        source_time = source.get('observed_at')
        time_basis = source.get('observed_at_basis', 'unknown')
        if config.get('observed_at_field'):
            raw_time = getpath(row, config['observed_at_field'])
            try:
                source_time = datetime.fromisoformat(str(raw_time).replace('Z', '+00:00')).isoformat()
                time_basis = f"source_field:{config['observed_at_field']}"
            except (TypeError, ValueError):
                pass
        if not source_time:
            errors.append({'source_record_id': current_record_id, 'reason': 'No source observation timestamp or documented metadata proxy'})
            continue
        envelope = {'source_id': source['source_id'], 'source_record_id': current_record_id,
                    'observed_at': source_time, 'ingested_at': ingested_at,
                    'licence': source.get('licence') or 'unknown',
                    'extractor_version': f'{VERSION}:{config_hash}',
                    'observed_at_basis': time_basis,
                    'provenance': {'resource_url': source.get('resource_url'), 'resource_sha256': source['resource_sha256'],
                                   'raw_record_sha256': raw_hash, 'row_number': index, 'config_sha256': config_hash}}
        for mapping in config['mappings']:
            field = mapping['canonical_field']
            if mapping['confidence'] < (0 if diagnostic else max(.8,config.get('minimum_field_confidence', .8))):
                continue
            raw_value = getpath(row, mapping['source_field'])
            try:
                value = transform(raw_value, mapping.get('transforms', []))
                if value is None:
                    continue
                values = value if isinstance(value, list) else [value]
                for item in values:
                    canonical = validate_value(item, field, ontology)
                    group, key = field.split('.')
                    observation = dict(envelope, canonical_field=field, **{group:{key:canonical}}, raw_value=raw_value,
                                       confidence={'source_reliability': source.get('source_reliability', .8),
                                                   'field_confidence': mapping['confidence']},
                                       derivation_level='L1' if mapping.get('transforms') else 'L0')
                    observation['provenance'] = dict(envelope['provenance'], source_field=mapping['source_field'])
                    if len(values) > 1 and mapping.get('separate_entities'):
                        observation['source_record_id'] = current_record_id+'#'+field+'='+canonical
                        observation['provenance']['parent_source_record_id'] = current_record_id
                        observation['provenance']['entity_scope'] = 'explicit_identifier_only_no_other_fields_assigned'
                    observation['observation_id'] = digest([source['source_id'], observation['source_record_id'], config_hash, field, canonical])
                    observations.append(observation)
            except (ValueError, TypeError, KeyError, IndexError) as exc:
                errors.append({'source_record_id': current_record_id, 'source_field': mapping['source_field'], 'canonical_field': field,
                               'raw_value': raw_value, 'reason': str(exc)})
    return observations, errors
